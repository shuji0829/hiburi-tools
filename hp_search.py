"""
医療機関ホームページ・メールアドレス検索スクリプト

Playwright + Google検索で事業者名+住所からHPを特定し、
HPからメールアドレスを抽出してExcelに反映する。

使い方:
  python hp_search.py --input 厚生局_新規指定一覧_R0704-R0803.xlsx --limit 100
  python hp_search.py --input 厚生局_新規指定一覧_R0704-R0803.xlsx --resume

前提:
  pip install playwright openpyxl
  playwright install chromium
"""

import argparse
import json
import os
import re
import sys
import time
import random

import openpyxl
from urllib.parse import urlparse

# 除外ドメイン（ポータルサイト等、公式HPではないもの）
EXCLUDE_DOMAINS = {
    "byoinnavi.jp", "caloo.jp", "fdoc.jp", "doctorsfile.jp",
    "medley.life", "hospita.jp", "myclinic.ne.jp", "scuel.me",
    "doctor-map.info", "epark.jp", "google.com", "google.co.jp",
    "yahoo.co.jp", "bing.com", "facebook.com", "twitter.com",
    "instagram.com", "youtube.com", "wikipedia.org",
    "itp.ne.jp", "mapion.co.jp", "navitime.co.jp",
    "ekiten.jp", "qlife.jp", "10man-doc.co.jp",
    "medinew.jp", "medimap.jp", "hospiten.jp",
    "kamponavi.com", "clinicnavi.jp", "mynavi.jp",
    "iryou.teikyouseido.mhlw.go.jp", "kurashi.yahoo.co.jp",
    "kusurinomadoguchi.com", "sokuyaku.jp", "clinics-app.com",
    "e-classa.net", "localplace.jp", "health.ne.jp",
    "job-medley.com", "m3.com", "pcareer.m3.com",
    "medicalnote.jp", "homemate-research-drugstore.com",
    "kaigokensaku.mhlw.go.jp", "jmap.jp", "opendata-japan.com",
    "shufoo.net", "tokubai.co.jp", "dpoint.docomo.ne.jp",
    "mapfan.com", "guppy.jp", "indeed.com", "stanby.co.jp",
    "haisha-yoyaku.jp", "dental-city.com",
}

# メールアドレスの正規表現
EMAIL_PATTERN = re.compile(
    r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}',
    re.IGNORECASE
)

# 除外メールアドレスパターン
EXCLUDE_EMAIL_PATTERNS = [
    r'.*@example\.com',
    r'.*@test\.com',
    r'.*\.png$',
    r'.*\.jpg$',
    r'.*\.gif$',
    r'.*\.svg$',
    r'.*wixpress\.com',
    r'.*sentry\.io',
    r'.*@sentry-next\.wixpress\.com',
]

# 進捗ファイル
PROGRESS_FILE = "hp_search_progress.json"


def is_official_domain(url):
    """ポータルサイトを除外して公式HPらしいか判定"""
    try:
        parsed = urlparse(url)
        domain = parsed.netloc.lower()
        if domain.startswith("www."):
            domain = domain[4:]
        for excl in EXCLUDE_DOMAINS:
            if domain == excl or domain.endswith("." + excl):
                return False
        return True
    except Exception:
        return False


def filter_email(email):
    """メールアドレスのフィルタリング"""
    for pattern in EXCLUDE_EMAIL_PATTERNS:
        if re.match(pattern, email, re.IGNORECASE):
            return False
    return True


def search_google_playwright(page, query, num_results=5):
    """Playwright を使ってGoogle検索し、結果URLを返す"""
    encoded_query = query.replace(" ", "+")
    url = f"https://www.google.co.jp/search?q={encoded_query}&hl=ja&gl=jp&num={num_results}"

    try:
        page.goto(url, timeout=15000, wait_until="domcontentloaded")
        # 少し待つ（レート制限対策）
        time.sleep(random.uniform(1.0, 2.0))

        # 検索結果のリンクを取得
        results = []
        # Google検索結果の主要セレクタ
        links = page.query_selector_all("div#search a[href]")
        for link in links:
            href = link.get_attribute("href")
            if href and href.startswith("http") and is_official_domain(href):
                # URLの正規化（トラッキングパラメータ除去）
                parsed = urlparse(href)
                clean_url = f"{parsed.scheme}://{parsed.netloc}{parsed.path}"
                if clean_url not in results:
                    results.append(clean_url)
                if len(results) >= num_results:
                    break
        return results
    except Exception as e:
        print(f"    検索エラー: {e}")
        return []


def extract_emails_from_page(page, url):
    """Playwrightを使ってページからメールアドレスを抽出"""
    try:
        page.goto(url, timeout=10000, wait_until="domcontentloaded")
        time.sleep(0.5)
        content = page.content()
        emails = EMAIL_PATTERN.findall(content)
        return [e for e in dict.fromkeys(emails) if filter_email(e)]
    except Exception:
        return []


def find_hp_and_email(page, name, address):
    """事業者名+住所から公式HPとメールアドレスを検索"""
    # 住所から都道府県+市区町村を抽出
    addr_short = ""
    addr_match = re.search(r'[^\d〒\s]{2,3}[都道府県].{2,6}[市区町村郡]', address or "")
    if addr_match:
        addr_short = addr_match.group()

    query = f"{name} {addr_short}"
    results = search_google_playwright(page, query, num_results=5)

    hp_url = None
    email = None

    if results:
        hp_url = results[0]  # 最初の公式サイト結果

        # HPからメール抽出
        parsed = urlparse(hp_url)
        top_url = f"{parsed.scheme}://{parsed.netloc}/"

        # まず検索結果のページ
        emails = extract_emails_from_page(page, hp_url)
        if emails:
            email = emails[0]

        # トップページ
        if not email and hp_url != top_url:
            emails = extract_emails_from_page(page, top_url)
            if emails:
                email = emails[0]

        # お問い合わせページ
        if not email:
            for suffix in ["/contact", "/contact/", "/about", "/access"]:
                contact_url = top_url.rstrip("/") + suffix
                emails = extract_emails_from_page(page, contact_url)
                if emails:
                    email = emails[0]
                    break

    return hp_url, email


def load_progress():
    """進捗ファイルを読み込み"""
    if os.path.exists(PROGRESS_FILE):
        with open(PROGRESS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_progress(progress):
    """進捗ファイルを保存"""
    with open(PROGRESS_FILE, "w", encoding="utf-8") as f:
        json.dump(progress, f, ensure_ascii=False, indent=2)


def main():
    parser = argparse.ArgumentParser(description="医療機関HP・メール検索 (Playwright+Google)")
    parser.add_argument("--input", required=True, help="入力Excelファイル")
    parser.add_argument("--sheet", default="新規指定一覧", help="シート名")
    parser.add_argument("--limit", type=int, default=0, help="処理件数上限 (0=全件)")
    parser.add_argument("--resume", action="store_true", help="前回の続きから再開")
    parser.add_argument("--headless", action="store_true", default=True,
                        help="ヘッドレスモード (デフォルト: True)")
    parser.add_argument("--no-headless", action="store_false", dest="headless",
                        help="ブラウザ表示モード")
    parser.add_argument("--delay", type=float, default=2.0,
                        help="検索間隔(秒) ※短すぎるとブロックされます")
    args = parser.parse_args()

    # Playwright インポート
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("エラー: Playwright がインストールされていません")
        print("  pip install playwright")
        print("  playwright install chromium")
        sys.exit(1)

    # Excel読み込み
    wb = openpyxl.load_workbook(args.input)
    ws = wb[args.sheet]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]

    name_idx = headers.index("事業者名")
    addr_idx = headers.index("住所")
    hp_idx = headers.index("ホームページ")
    email_idx = headers.index("メールアドレス")
    wb.close()

    # 進捗管理
    progress = load_progress() if args.resume else {}

    total = len(rows) - 1
    limit = args.limit if args.limit > 0 else total
    searched = 0
    found_hp = 0
    found_email = 0

    print(f"=== HP・メール検索 (Playwright + Google) ===")
    print(f"対象: {total} 件 (上限: {limit} 件)")
    print(f"既処理: {len(progress)} 件")
    print(f"ヘッドレス: {args.headless}")
    print()

    # 書き込み用にExcelを再オープン
    wb = openpyxl.load_workbook(args.input)
    ws = wb[args.sheet]

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=args.headless)
        context = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            locale="ja-JP",
        )
        page = context.new_page()

        try:
            for row_num, row in enumerate(rows[1:], start=2):
                if searched >= limit:
                    break

                name = str(row[name_idx] or "").strip()
                address = str(row[addr_idx] or "").strip()

                if not name:
                    continue

                # 既にHPが入力済みならスキップ
                existing_hp = row[hp_idx]
                if existing_hp:
                    continue

                # 進捗チェック（前回の結果を書き込み）
                if name in progress:
                    hp = progress[name].get("hp")
                    em = progress[name].get("email")
                    if hp:
                        ws.cell(row=row_num, column=hp_idx + 1, value=hp)
                    if em:
                        ws.cell(row=row_num, column=email_idx + 1, value=em)
                    continue

                # Google検索
                print(f"  [{searched + 1}/{limit}] {name}...", end=" ", flush=True)
                try:
                    hp, em = find_hp_and_email(page, name, address)
                except Exception as e:
                    print(f"エラー: {e}")
                    progress[name] = {"hp": None, "email": None}
                    searched += 1
                    continue

                # 結果保存
                progress[name] = {"hp": hp, "email": em}

                if hp:
                    ws.cell(row=row_num, column=hp_idx + 1, value=hp)
                    found_hp += 1
                if em:
                    ws.cell(row=row_num, column=email_idx + 1, value=em)
                    found_email += 1

                status = f"HP: {hp or 'なし'}"
                if em:
                    status += f" | Email: {em}"
                print(status)

                searched += 1

                # 10件ごとに中間保存
                if searched % 10 == 0:
                    print(f"\n  --- {searched}/{limit} 件処理 (HP: {found_hp}, Email: {found_email}) ---\n")
                    save_progress(progress)
                    wb.save(args.input)

                # レート制限対策
                time.sleep(args.delay + random.uniform(0, 1.0))

        except KeyboardInterrupt:
            print("\n中断されました。保存中...")
        finally:
            browser.close()

    # 最終保存
    save_progress(progress)
    wb.save(args.input)
    wb.close()

    print(f"\n=== 完了 ===")
    print(f"  処理件数: {searched}")
    print(f"  HP発見: {found_hp}")
    print(f"  メール発見: {found_email}")
    print(f"  進捗保存: {PROGRESS_FILE}")


if __name__ == "__main__":
    main()
