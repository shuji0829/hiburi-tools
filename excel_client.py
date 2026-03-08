"""
Excel 読み書きクライアント
"""

import openpyxl
from config import COLUMN_ALIASES


def find_header_row(rows):
    """ヘッダー行を自動検出（施設名・メールアドレス等が含まれる行）"""
    search_keys = set(COLUMN_ALIASES.keys())
    for idx, row in enumerate(rows):
        cells = [str(c).strip() if c else "" for c in row]
        matched = [c for c in cells if c in search_keys]
        if len(matched) >= 3:
            return idx
    return None


def read_excel(file_path, sheet_name, header_row=None):
    """Excelファイルを読み込み、マッピング対象列のみ返す"""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("Excelファイルにデータがありません")
        wb.close()
        return []

    if header_row is not None:
        h_idx = header_row - 1
    else:
        h_idx = find_header_row(rows)
        if h_idx is None:
            print("エラー: ヘッダー行が見つかりません。--header-row で指定してください")
            print("  先頭5行の内容:")
            for i, r in enumerate(rows[:5], 1):
                print(f"    行{i}: {r}")
            wb.close()
            return []
        print(f"  ヘッダー行を自動検出: {h_idx + 1}行目")

    headers = [str(h).strip() if h else "" for h in rows[h_idx]]
    allowed_columns = set(COLUMN_ALIASES.keys())
    matched = [h for h in headers if h in allowed_columns]
    ignored = [h for h in headers if h and h not in allowed_columns]
    print(f"  マッピング対象列: {matched}")
    if ignored:
        print(f"  無視する列: {ignored}")

    data = []
    for row in rows[h_idx + 1:]:
        row_dict = {}
        for i, header in enumerate(headers):
            if header in allowed_columns and i < len(row):
                row_dict[header] = row[i]
        if any(v is not None for v in row_dict.values()):
            data.append(row_dict)

    wb.close()
    return data


def write_excel(file_path, sheet_name, rows_data, headers=None):
    """データをExcelファイルに書き出す（新規作成または上書き）"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    if not rows_data:
        wb.save(file_path)
        wb.close()
        return

    if headers is None:
        headers = list(rows_data[0].keys())
        # _page_id等の内部キーを除外
        headers = [h for h in headers if not h.startswith("_")]

    # ヘッダー行
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    # データ行
    for row_idx, row_data in enumerate(rows_data, 2):
        for col, header in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col, value=row_data.get(header, ""))

    wb.save(file_path)
    wb.close()
    print(f"  Excelファイル保存: {file_path} ({len(rows_data)} 行)")


def update_excel_in_place(file_path, sheet_name, updates, key_column="事業者名"):
    """既存Excelファイルの行をキー列で照合して更新"""
    wb = openpyxl.load_workbook(file_path)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        wb.close()
        return 0

    h_idx = find_header_row(rows)
    if h_idx is None:
        wb.close()
        return 0

    headers = [str(h).strip() if h else "" for h in rows[h_idx]]

    # キー列のインデックス
    key_col_idx = None
    for i, h in enumerate(headers):
        if h == key_column or COLUMN_ALIASES.get(h) == key_column:
            key_col_idx = i
            break
    if key_col_idx is None:
        wb.close()
        return 0

    # updates を辞書化 {key_value: {col: val}}
    update_map = {}
    for u in updates:
        key_val = u.get(key_column)
        if key_val:
            update_map[key_val] = u

    updated_count = 0
    for row_num in range(h_idx + 2, ws.max_row + 1):  # 1-indexed, skip header
        cell_val = ws.cell(row=row_num, column=key_col_idx + 1).value
        if cell_val and str(cell_val).strip() in update_map:
            upd = update_map[str(cell_val).strip()]
            for col_idx, header in enumerate(headers):
                notion_prop = COLUMN_ALIASES.get(header, header)
                if notion_prop in upd and upd[notion_prop] is not None:
                    ws.cell(row=row_num, column=col_idx + 1, value=upd[notion_prop])
            updated_count += 1

    wb.save(file_path)
    wb.close()
    return updated_count
